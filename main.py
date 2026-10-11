from collections import defaultdict
from pathlib import Path

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet


INPUT_FILE = Path("inventory.xlsx")
OUTPUT_FILE = Path("inventory_with_total_value.xlsx")
SHEET_NAME = "Sheet1"

PRODUCT_NUMBER_COLUMN = 1
INVENTORY_COLUMN = 2
PRICE_COLUMN = 3
SUPPLIER_COLUMN = 4
TOTAL_VALUE_COLUMN = 5

LOW_INVENTORY_THRESHOLD = 10


def analyze_inventory(
    worksheet: Worksheet,
) -> tuple[dict[str, int], dict[str, float], dict[int, int]]:
    products_per_supplier: dict[str, int] = defaultdict(int)
    total_value_per_supplier: dict[str, float] = defaultdict(float)
    products_under_threshold: dict[int, int] = {}

    for row_number in range(2, worksheet.max_row + 1):
        product_number = worksheet.cell(row=row_number, column=PRODUCT_NUMBER_COLUMN).value
        inventory = worksheet.cell(row=row_number, column=INVENTORY_COLUMN).value
        price = worksheet.cell(row=row_number, column=PRICE_COLUMN).value
        supplier_name = worksheet.cell(row=row_number, column=SUPPLIER_COLUMN).value

        if supplier_name is None:
            raise ValueError(
                f"Missing supplier name in row {row_number}"
            )

        if inventory is None or price is None:
            raise ValueError(
                f"Missing inventory or price in row {row_number}"
            )

        inventory = int(inventory)
        price = float(price)

        products_per_supplier[supplier_name] += 1

        inventory_value = inventory * price
        total_value_per_supplier[supplier_name] += inventory_value

        if inventory < LOW_INVENTORY_THRESHOLD and product_number is not None:
            products_under_threshold[int(product_number)] = inventory

        worksheet.cell(
            row=row_number,
            column=TOTAL_VALUE_COLUMN,
            value=inventory_value,
        )

    return (
        dict(products_per_supplier),
        dict(total_value_per_supplier),
        products_under_threshold,
    )


def main() -> None:
    if not INPUT_FILE.exists():
        raise FileNotFoundError(
            f"Inventory file does not exist: {INPUT_FILE}"
        )

    workbook = openpyxl.load_workbook(INPUT_FILE)

    if SHEET_NAME not in workbook.sheetnames:
        raise ValueError(
            f"Worksheet '{SHEET_NAME}' not found. "
            f"Available sheets: {workbook.sheetnames}"
        )

    worksheet = workbook[SHEET_NAME]

    (
        products_per_supplier,
        total_value_per_supplier,
        products_under_threshold,
    ) = analyze_inventory(worksheet)

    print("Products per supplier:")
    print(products_per_supplier)

    print("\nTotal inventory value per supplier:")
    print(total_value_per_supplier)

    print(f"\nProducts with inventory below {LOW_INVENTORY_THRESHOLD}:")
    print(products_under_threshold)

    workbook.save(OUTPUT_FILE)

    print(f"\nUpdated workbook saved to: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()